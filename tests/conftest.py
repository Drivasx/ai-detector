"""Fixtures compartidas: generan documentos de prueba en directorios temporales."""

import pytest
from docx import Document


def _escape_pdf_string(text: str) -> bytes:
    """Codifica una línea como cadena literal PDF (WinAnsi ~ cp1252)."""
    raw = text.encode("cp1252")
    return raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


def build_pdf(pages: list[list[str]]) -> bytes:
    """Construye un PDF mínimo válido con una línea de texto por elemento de cada página."""
    page_count = len(pages)
    font_id = 3 + 2 * page_count
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids ["
        + b" ".join(f"{3 + 2 * i} 0 R".encode() for i in range(page_count))
        + f"] /Count {page_count} >>".encode(),
    ]
    for i, lines in enumerate(pages):
        content_id = 4 + 2 * i
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>".encode()
        )
        stream = b"BT /F1 12 Tf 14 TL 72 720 Td " + b"".join(
            b"(" + _escape_pdf_string(line) + b") Tj T* " for line in lines
        ) + b"ET"
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_offset = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode()
    return bytes(out)


@pytest.fixture
def make_pdf(tmp_path):
    """Escribe un PDF de prueba y devuelve su ruta."""

    def _make(pages: list[list[str]], name: str = "doc.pdf"):
        path = tmp_path / name
        path.write_bytes(build_pdf(pages))
        return path

    return _make


@pytest.fixture
def make_docx(tmp_path):
    """Escribe un .docx con un párrafo por elemento y devuelve su ruta."""

    def _make(paragraphs: list[str], name: str = "doc.docx", table_rows: list[list[str]] | None = None):
        document = Document()
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
        if table_rows:
            table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
            for row, values in zip(table.rows, table_rows):
                for cell, value in zip(row.cells, values):
                    cell.text = value
        path = tmp_path / name
        document.save(path)
        return path

    return _make
